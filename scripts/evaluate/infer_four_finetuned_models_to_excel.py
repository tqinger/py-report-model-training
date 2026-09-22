"""Run four Qwen3-4B LoRA adapters on four held-out cases each.

The script writes the full prompt (system/user messages) and the generated answer
to a single Excel workbook.  It deliberately uses the supplied validation splits
for the JSONL-trained adapters and the r10 held-out split for the tongue adapter.
"""

from __future__ import annotations

import argparse
import gc
import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import torch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from peft import PeftModel
from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig

from tcm_qwen_eval.dataset import load_examples
from tcm_qwen_eval.tongue_qlora import split_tongue_examples

DEFAULT_BASE_MODEL = Path(
    "artifacts/hf_cache/hub/models--Qwen--Qwen3-4B/snapshots/"
    "1cfa9a7208912126459214e8b04321603b3df60c"
)
DEFAULT_OUTPUT = Path(
    "outputs/four_finetuned_models_4cases_20260827/"
    "four_finetuned_models_4cases.xlsx"
)


@dataclass(frozen=True)
class ModelSpec:
    name: str
    sheet_name: str
    adapter: Path
    selection_source: str
    data_path: Path
    selection_kind: str


MODEL_SPECS = (
    ModelSpec(
        name="全科中医报告模型",
        sheet_name="全科报告",
        adapter=Path("artifacts/qwen3-4b-holistic-50k-qlora-windows"),
        selection_source="data/holistic_50k/val.jsonl（验证集）",
        data_path=Path("data/holistic_50k/val.jsonl"),
        selection_kind="jsonl_validation",
    ),
    ModelSpec(
        name="病例润色模型",
        sheet_name="病例润色",
        adapter=Path("artifacts/qwen3-4b-case-polish-lora-windows"),
        selection_source="data/case-polish-lora/val.jsonl（验证集）",
        data_path=Path("data/case-polish-lora/val.jsonl"),
        selection_kind="jsonl_validation",
    ),
    ModelSpec(
        name="五态调养模型",
        sheet_name="五态调养",
        adapter=Path("artifacts/qwen3-4b-wutai-20-qlora-windows"),
        selection_source="data/wutai_20/val.jsonl（验证集）",
        data_path=Path("data/wutai_20/val.jsonl"),
        selection_kind="jsonl_validation",
    ),
    ModelSpec(
        name="舌象分析模型",
        sheet_name="舌象分析",
        adapter=Path("artifacts/qwen3-4b-tongue-conversations-qlora"),
        selection_source="data/conversations 的 r10 留出集",
        data_path=Path("data/conversations"),
        selection_kind="tongue_test",
    ),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run four fine-tuned Qwen3-4B adapters on four cases and export Excel."
    )
    parser.add_argument("--base-model", type=Path, default=DEFAULT_BASE_MODEL)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--cases-per-model", type=int, default=4)
    parser.add_argument("--max-new-tokens", type=int, default=768)
    return parser.parse_args()


def validate_messages(messages: object, source: str) -> list[dict[str, str]]:
    if not isinstance(messages, list) or len(messages) < 2:
        raise ValueError(f"{source}: expected at least a prompt and an assistant response")
    normalized: list[dict[str, str]] = []
    for message in messages:
        if not isinstance(message, dict) or "role" not in message or "content" not in message:
            raise ValueError(f"{source}: invalid chat message")
        normalized.append({"role": str(message["role"]), "content": str(message["content"])})
    if normalized[-1]["role"] != "assistant":
        raise ValueError(f"{source}: final message must be the reference assistant response")
    return normalized


def select_jsonl_validation_cases(path: Path, count: int) -> list[tuple[str, list[dict[str, str]]]]:
    """Return the earliest prompt-distinct records from an existing validation split."""
    selected: list[tuple[str, list[dict[str, str]]]] = []
    seen_prompts: set[tuple[tuple[str, str], ...]] = set()
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            record = json.loads(line)
            messages = validate_messages(record.get("messages"), f"{path}:{line_number}")
            prompt_messages = messages[:-1]
            signature = tuple((item["role"], item["content"]) for item in prompt_messages)
            if signature in seen_prompts:
                continue
            seen_prompts.add(signature)
            selected.append((f"验证集第 {line_number} 条", messages))
            if len(selected) == count:
                return selected
    raise ValueError(f"{path}: found only {len(selected)} prompt-distinct validation cases; need {count}")


def select_tongue_test_cases(data_path: Path, count: int) -> list[tuple[str, list[dict[str, str]]]]:
    test_examples = split_tongue_examples(load_examples(data_path))["test"]
    if len(test_examples) < count:
        raise ValueError(f"{data_path}: test split has only {len(test_examples)} cases; need {count}")
    return [(example.id, example.messages) for example in sorted(test_examples, key=lambda item: item.id)[:count]]


def select_cases(spec: ModelSpec, count: int) -> list[tuple[str, list[dict[str, str]]]]:
    if spec.selection_kind == "jsonl_validation":
        return select_jsonl_validation_cases(spec.data_path, count)
    if spec.selection_kind == "tongue_test":
        return select_tongue_test_cases(spec.data_path, count)
    raise ValueError(f"Unsupported selection kind: {spec.selection_kind}")


def render_input(messages: list[dict[str, str]]) -> str:
    return "\n\n".join(
        f"[{message['role']}]\n{message['content']}" for message in messages[:-1]
    )


def resolve_adapter_path(adapter_root: Path) -> Path:
    """Resolve a saved adapter root or its sole Trainer checkpoint directory."""
    if (adapter_root / "adapter_config.json").is_file():
        return adapter_root
    checkpoints = sorted(
        path
        for path in adapter_root.glob("checkpoint-*")
        if (path / "adapter_config.json").is_file()
    )
    if len(checkpoints) == 1:
        return checkpoints[0]
    if not checkpoints:
        raise FileNotFoundError(f"No adapter_config.json found in {adapter_root} or its checkpoints")
    raise ValueError(
        f"{adapter_root} has multiple adapter checkpoints; select one explicitly: {checkpoints}"
    )


def load_adapter(base_model_path: Path, adapter_root: Path) -> tuple[Any, Any, Path]:
    adapter_path = resolve_adapter_path(adapter_root)
    quantization = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_use_double_quant=True,
        bnb_4bit_compute_dtype=torch.bfloat16,
    )
    tokenizer_source = adapter_path if (adapter_path / "tokenizer.json").is_file() else base_model_path
    tokenizer = AutoTokenizer.from_pretrained(tokenizer_source, local_files_only=True)
    if tokenizer.pad_token_id is None:
        tokenizer.pad_token = tokenizer.eos_token
    tokenizer.padding_side = "right"
    base = AutoModelForCausalLM.from_pretrained(
        base_model_path,
        quantization_config=quantization,
        torch_dtype=torch.bfloat16,
        device_map={"": 0},
        low_cpu_mem_usage=True,
        local_files_only=True,
    )
    model = PeftModel.from_pretrained(base, adapter_path, local_files_only=True)
    model.generation_config.do_sample = False
    model.generation_config.temperature = None
    model.generation_config.top_p = None
    model.generation_config.top_k = None
    model.eval()
    return model, tokenizer, adapter_path


@torch.inference_mode()
def generate(
    model: Any, tokenizer: Any, messages: list[dict[str, str]], max_new_tokens: int
) -> tuple[str, int, bool]:
    try:
        inputs = tokenizer.apply_chat_template(
            messages[:-1],
            tokenize=True,
            add_generation_prompt=True,
            enable_thinking=False,
            return_dict=True,
            return_tensors="pt",
        )
    except TypeError:
        inputs = tokenizer.apply_chat_template(
            messages[:-1],
            tokenize=True,
            add_generation_prompt=True,
            return_dict=True,
            return_tensors="pt",
        )
    inputs = inputs.to(model.device)
    generated = model.generate(
        **inputs,
        do_sample=False,
        max_new_tokens=max_new_tokens,
        pad_token_id=tokenizer.pad_token_id,
        eos_token_id=tokenizer.eos_token_id,
    )
    output_ids = generated[0, inputs.input_ids.shape[1] :]
    output_token_count = int(output_ids.shape[0])
    hit_token_limit = output_token_count >= max_new_tokens
    return tokenizer.decode(output_ids, skip_special_tokens=True).strip(), output_token_count, hit_token_limit


def style_sheet(sheet: Any, widths: tuple[int, ...], row_height: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="thin", color="D1D5DB"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment
            cell.border = border
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = row_height


def create_workbook(args: argparse.Namespace) -> tuple[Workbook, Any]:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "说明"
    overview.append(["项目", "内容"])
    overview.append(["测试范围", f"4 个微调模型，每个模型 {args.cases_per_model} 个案例"])
    overview.append(["生成方式", "贪心生成（do_sample=False），已关闭 Qwen3 思考模式"])
    overview.append(["最大生成长度", f"{args.max_new_tokens} tokens"])
    overview.append(["基座模型", str(args.base_model)])
    overview.append(["模型加载方式", "NF4 4-bit QLoRA，bfloat16 计算"])
    overview.append(["案例选择", "前三个模型取验证集中的最早 4 条输入不同记录；舌象模型取 r10 留出集排序后的前 4 条。"])
    overview.append(["模型", "适配器路径"])
    for spec in MODEL_SPECS:
        overview.append([spec.name, str(spec.adapter)])
    style_sheet(overview, (24, 108), 38)
    return workbook, overview


def add_model_sheet(
    workbook: Workbook,
    spec: ModelSpec,
    rows: list[dict[str, object]],
) -> None:
    sheet = workbook.create_sheet(spec.sheet_name)
    sheet.append(
        [
            "模型",
            "案例",
            "案例来源",
            "输入（system / user）",
            "模型输出",
            "输出 tokens",
            "耗时（秒）",
            "达到长度上限",
        ]
    )
    for row in rows:
        sheet.append(
            [
                spec.name,
                row["case_id"],
                spec.selection_source,
                row["input"],
                row["output"],
                row["output_tokens"],
                row["latency_seconds"],
                "是" if row["hit_token_limit"] else "否",
            ]
        )
    style_sheet(sheet, (18, 38, 44, 72, 90, 14, 14, 16), 420)


def main() -> None:
    args = parse_args()
    if args.cases_per_model <= 0:
        raise ValueError("--cases-per-model must be positive")
    if args.max_new_tokens <= 0:
        raise ValueError("--max-new-tokens must be positive")
    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is unavailable; 4-bit QLoRA inference requires a CUDA GPU.")
    if not args.base_model.is_dir():
        raise FileNotFoundError(f"Base model not found: {args.base_model}")
    for spec in MODEL_SPECS:
        if not spec.adapter.is_dir():
            raise FileNotFoundError(f"Adapter directory not found: {spec.adapter}")

    workbook, _ = create_workbook(args)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    for model_index, spec in enumerate(MODEL_SPECS, start=1):
        cases = select_cases(spec, args.cases_per_model)
        print(f"[{model_index}/{len(MODEL_SPECS)}] Loading {spec.name}...", flush=True)
        model, tokenizer, resolved_adapter = load_adapter(args.base_model, spec.adapter)
        print(f"  Adapter resolved to: {resolved_adapter}", flush=True)
        rows: list[dict[str, object]] = []
        try:
            for case_index, (case_id, messages) in enumerate(cases, start=1):
                started = time.perf_counter()
                output, output_tokens, hit_token_limit = generate(
                    model, tokenizer, messages, args.max_new_tokens
                )
                elapsed = round(time.perf_counter() - started, 3)
                rows.append(
                    {
                        "case_id": case_id,
                        "input": render_input(messages),
                        "output": output,
                        "output_tokens": output_tokens,
                        "latency_seconds": elapsed,
                        "hit_token_limit": hit_token_limit,
                    }
                )
                print(
                    f"  case {case_index}/{len(cases)} complete: {case_id} ({elapsed:.3f}s)",
                    flush=True,
                )
        finally:
            del model
            del tokenizer
            gc.collect()
            torch.cuda.empty_cache()

        add_model_sheet(workbook, spec, rows)
        workbook.save(args.output)
        print(f"  Workbook checkpoint saved: {args.output}", flush=True)

    print(f"Completed: {args.output}", flush=True)


if __name__ == "__main__":
    main()
