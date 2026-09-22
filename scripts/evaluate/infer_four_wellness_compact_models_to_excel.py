"""Generate five held-out wellness cases for each compact Qwen3-4B LoRA adapter.

The workbook is deliberately written with openpyxl so reviewers can inspect the
prompt, generated answer, and generation metadata in a single .xlsx file.
"""

from __future__ import annotations

import argparse
import gc
import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import torch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from peft import PeftModel
from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig

DEFAULT_BASE_MODEL = Path(
    "artifacts/hf_cache/hub/models--Qwen--Qwen3-4B/snapshots/"
    "1cfa9a7208912126459214e8b04321603b3df60c"
)
DEFAULT_OUTPUT = Path("artifacts/wellness_compact_four_models_5cases_sampling_20260901.xlsx")
AGE_GENDER_ADVICE_CONSTRAINT = (
    "重要限定：必须同时考虑用户提供的年龄和性别，仅输出与该年龄和性别相匹配、"
    "可执行且安全的建议；不适用时应省略，不得给出明显不适配的通用建议。"
)


@dataclass(frozen=True)
class ModelSpec:
    key: str
    name: str
    sheet_name: str
    adapter: Path
    validation_data: Path
    apply_age_gender_constraint: bool = False


MODEL_SPECS = (
    ModelSpec(
        key="holistic",
        name="整体康养模型",
        sheet_name="整体康养",
        adapter=Path("artifacts/qwen3-4b-wellness-compact-holistic-qlora-windows"),
        validation_data=Path("data/wellness_lora_compact/export/holistic/val.jsonl"),
        apply_age_gender_constraint=True,
    ),
    ModelSpec(
        key="tongue",
        name="舌象康养模型",
        sheet_name="舌象康养",
        adapter=Path("artifacts/qwen3-4b-wellness-compact-tongue-qlora-windows"),
        validation_data=Path("data/wellness_lora_compact/export/tongue/val.jsonl"),
        apply_age_gender_constraint=True,
    ),
    ModelSpec(
        key="tongue_constitution",
        name="舌象体质模型",
        sheet_name="舌象体质",
        adapter=Path(
            "artifacts/qwen3-4b-wellness-compact-tongue-constitution-qlora-windows"
        ),
        validation_data=Path(
            "data/wellness_lora_compact/export/tongue_constitution/val.jsonl"
        ),
        apply_age_gender_constraint=True,
    ),
    ModelSpec(
        key="wutai",
        name="五态调养模型",
        sheet_name="五态调养",
        adapter=Path("artifacts/qwen3-4b-wellness-compact-wutai-qlora-windows"),
        validation_data=Path("data/wellness_lora_compact/export/wutai/val.jsonl"),
    ),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate held-out cases for four compact wellness LoRA adapters."
    )
    parser.add_argument("--base-model", type=Path, default=DEFAULT_BASE_MODEL)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--cases-per-model", type=int, default=5)
    parser.add_argument("--max-new-tokens", type=int, default=2048)
    parser.add_argument("--temperature", type=float, default=0.7)
    parser.add_argument("--top-p", type=float, default=0.8)
    parser.add_argument("--seed", type=int, default=20260901)
    parser.add_argument(
        "--only",
        choices=[spec.key for spec in MODEL_SPECS],
        nargs="+",
        help="Generate only the named model(s). Use --resume to preserve other workbook sheets.",
    )
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Load --output and replace only the generated model sheet(s).",
    )
    return parser.parse_args()


def validate_messages(messages: object, source: str) -> list[dict[str, str]]:
    if not isinstance(messages, list) or len(messages) < 2:
        raise ValueError(f"{source}: expected a prompt and an assistant reference answer")

    normalized: list[dict[str, str]] = []
    for message in messages:
        if not isinstance(message, dict) or {"role", "content"} - message.keys():
            raise ValueError(f"{source}: invalid chat message")
        normalized.append({"role": str(message["role"]), "content": str(message["content"])})

    if normalized[-1]["role"] != "assistant":
        raise ValueError(f"{source}: final message must be the reference assistant response")
    return normalized


def select_validation_cases(
    path: Path, count: int
) -> list[tuple[str, int, list[dict[str, str]]]]:
    """Select earliest prompt-distinct examples from the supplied validation split."""
    selected: list[tuple[str, int, list[dict[str, str]]]] = []
    seen_prompts: set[tuple[tuple[str, str], ...]] = set()
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            record = json.loads(line)
            messages = validate_messages(record.get("messages"), f"{path}:{line_number}")
            prompt_messages = messages[:-1]
            signature = tuple((item["role"], item["content"]) for item in prompt_messages)
            if signature in seen_prompts:
                continue
            seen_prompts.add(signature)
            case_id = str(record.get("id", f"validation line {line_number}"))
            selected.append((case_id, line_number, messages))
            if len(selected) == count:
                return selected

    raise ValueError(f"{path}: found only {len(selected)} distinct prompts; need {count}")


def render_input(messages: list[dict[str, str]]) -> str:
    return "\n\n".join(
        f"[{message['role']}]\n{message['content']}" for message in messages[:-1]
    )


def apply_age_gender_constraint(messages: list[dict[str, str]]) -> list[dict[str, str]]:
    """Append the reviewer-requested advice constraint to the system prompt."""
    constrained_messages = [message.copy() for message in messages]
    for message in constrained_messages[:-1]:
        if message["role"] == "system":
            message["content"] = f"{message['content']}\n\n{AGE_GENDER_ADVICE_CONSTRAINT}"
            return constrained_messages
    constrained_messages.insert(0, {"role": "system", "content": AGE_GENDER_ADVICE_CONSTRAINT})
    return constrained_messages


def resolve_adapter_path(adapter_root: Path) -> Path:
    if (adapter_root / "adapter_config.json").is_file():
        return adapter_root
    checkpoints = sorted(
        path
        for path in adapter_root.glob("checkpoint-*")
        if (path / "adapter_config.json").is_file()
    )
    if len(checkpoints) == 1:
        return checkpoints[0]
    if not checkpoints:
        raise FileNotFoundError(f"No adapter_config.json found in {adapter_root}")
    raise ValueError(f"{adapter_root} has multiple valid adapter checkpoints: {checkpoints}")


def load_adapter(base_model_path: Path, adapter_root: Path) -> tuple[Any, Any, Path]:
    adapter_path = resolve_adapter_path(adapter_root)
    quantization = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_use_double_quant=True,
        bnb_4bit_compute_dtype=torch.bfloat16,
    )
    tokenizer = AutoTokenizer.from_pretrained(adapter_path, local_files_only=True)
    if tokenizer.pad_token_id is None:
        tokenizer.pad_token = tokenizer.eos_token
    tokenizer.padding_side = "right"
    base = AutoModelForCausalLM.from_pretrained(
        base_model_path,
        quantization_config=quantization,
        torch_dtype=torch.bfloat16,
        device_map={"": 0},
        low_cpu_mem_usage=True,
        local_files_only=True,
    )
    model = PeftModel.from_pretrained(base, adapter_path, local_files_only=True)
    model.eval()
    return model, tokenizer, adapter_path


@torch.inference_mode()
def generate(
    model: Any,
    tokenizer: Any,
    messages: list[dict[str, str]],
    max_new_tokens: int,
    temperature: float,
    top_p: float,
) -> tuple[str, int, bool]:
    inputs = tokenizer.apply_chat_template(
        messages[:-1],
        tokenize=True,
        add_generation_prompt=True,
        enable_thinking=False,
        return_dict=True,
        return_tensors="pt",
    ).to(model.device)
    generated = model.generate(
        **inputs,
        do_sample=True,
        temperature=temperature,
        top_p=top_p,
        max_new_tokens=max_new_tokens,
        pad_token_id=tokenizer.pad_token_id,
        eos_token_id=tokenizer.eos_token_id,
    )
    output_ids = generated[0, inputs.input_ids.shape[1] :]
    output_token_count = int(output_ids.shape[0])
    return (
        tokenizer.decode(output_ids, skip_special_tokens=True).strip(),
        output_token_count,
        output_token_count >= max_new_tokens,
    )


def style_sheet(sheet: Any, widths: tuple[int, ...], row_height: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="thin", color="D1D5DB"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment
            cell.border = border
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = row_height


def create_workbook(args: argparse.Namespace) -> Workbook:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "说明"
    overview.append(["项目", "内容"])
    overview.append(["测试范围", f"4 个康养模型，每个模型 {args.cases_per_model} 个验证集案例"])
    overview.append(["案例来源", "各模型训练使用的 val.jsonl；选取最早的 5 个输入不重复案例"])
    overview.append(
        [
            "生成方式",
            (
                "采样生成（do_sample=True），"
                f"temperature={args.temperature}，top_p={args.top_p}，"
                f"seed={args.seed}；已关闭 Qwen3 思考模式"
            ),
        ]
    )
    overview.append(["最大生成长度", f"{args.max_new_tokens} tokens"])
    overview.append(["基座模型", str(args.base_model)])
    overview.append(["加载方式", "NF4 4-bit QLoRA，bfloat16 计算"])
    overview.append(["提示词限定", AGE_GENDER_ADVICE_CONSTRAINT])
    overview.append(["模型", "适配器路径"])
    for spec in MODEL_SPECS:
        overview.append([spec.name, str(spec.adapter)])
    style_sheet(overview, (22, 118), 38)
    return workbook


def add_model_sheet(
    workbook: Workbook,
    spec: ModelSpec,
    adapter_path: Path,
    rows: list[dict[str, object]],
    index: int | None = None,
) -> None:
    sheet = workbook.create_sheet(spec.sheet_name, index=index)
    sheet.append(
        [
            "模型",
            "案例编号",
            "验证集行号",
            "案例来源",
            "输入（system / user）",
            "模型生成答案",
            "输出 tokens",
            "耗时（秒）",
            "达到长度上限",
            "实际适配器路径",
        ]
    )
    for row in rows:
        sheet.append(
            [
                spec.name,
                row["case_id"],
                row["line_number"],
                str(spec.validation_data),
                row["input"],
                row["output"],
                row["output_tokens"],
                row["latency_seconds"],
                "是" if row["hit_token_limit"] else "否",
                str(adapter_path),
            ]
        )
    style_sheet(sheet, (18, 32, 12, 54, 72, 90, 14, 14, 16, 58), 420)
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 7).number_format = "#,##0"
        sheet.cell(row_index, 8).number_format = "0.000"


def add_resume_sampling_note(workbook: Workbook, spec: ModelSpec, args: argparse.Namespace) -> None:
    """Record exceptional sampling parameters when one worksheet is regenerated."""
    overview = workbook["说明"]
    label = f"{spec.name}（重生成采样参数）"
    value = (
        "do_sample=True，"
        f"temperature={args.temperature}，top_p={args.top_p}，"
        f"seed={args.seed + MODEL_SPECS.index(spec)}；"
        + ("已启用年龄/性别适配建议限定" if spec.apply_age_gender_constraint else "未修改提示词")
    )
    for row_index in range(2, overview.max_row + 1):
        if overview.cell(row_index, 1).value == label:
            overview.cell(row_index, 2).value = value
            break
    else:
        overview.append([label, value])
    style_sheet(overview, (22, 118), 38)


def verify_workbook(
    output: Path, expected_cases: int, expected_specs: tuple[ModelSpec, ...]
) -> None:
    workbook = load_workbook(output, read_only=True, data_only=False)
    expected_sheet_names = ["说明", *(spec.sheet_name for spec in expected_specs)]
    if workbook.sheetnames != expected_sheet_names:
        raise ValueError(f"Unexpected worksheets: {workbook.sheetnames}")
    for spec in expected_specs:
        sheet = workbook[spec.sheet_name]
        if sheet.max_row != expected_cases + 1:
            raise ValueError(f"{spec.name}: expected {expected_cases} generated rows")
        generated_answers = [sheet.cell(row, 6).value for row in range(2, sheet.max_row + 1)]
        if any(not isinstance(answer, str) or not answer.strip() for answer in generated_answers):
            raise ValueError(f"{spec.name}: workbook contains a blank generated answer")
    workbook.close()


def main() -> None:
    args = parse_args()
    if args.cases_per_model <= 0:
        raise ValueError("--cases-per-model must be positive")
    if args.max_new_tokens <= 0:
        raise ValueError("--max-new-tokens must be positive")
    if args.temperature <= 0:
        raise ValueError("--temperature must be positive when sampling")
    if not 0 < args.top_p <= 1:
        raise ValueError("--top-p must be in (0, 1]")
    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is unavailable; 4-bit QLoRA inference requires a CUDA GPU.")
    if not args.base_model.is_dir():
        raise FileNotFoundError(f"Base model not found: {args.base_model}")
    for spec in MODEL_SPECS:
        if not spec.adapter.is_dir():
            raise FileNotFoundError(f"Adapter directory not found: {spec.adapter}")
        if not spec.validation_data.is_file():
            raise FileNotFoundError(f"Validation data not found: {spec.validation_data}")

    selected_specs = (
        tuple(spec for spec in MODEL_SPECS if spec.key in args.only) if args.only else MODEL_SPECS
    )
    if args.resume:
        if not args.output.is_file():
            raise FileNotFoundError(f"Cannot resume; workbook not found: {args.output}")
        workbook = load_workbook(args.output)
    else:
        workbook = create_workbook(args)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    for model_index, spec in enumerate(selected_specs, start=1):
        cases = select_validation_cases(spec.validation_data, args.cases_per_model)
        print(f"[{model_index}/{len(selected_specs)}] Loading {spec.name}...", flush=True)
        model, tokenizer, resolved_adapter = load_adapter(args.base_model, spec.adapter)
        print(f"  Adapter resolved to: {resolved_adapter}", flush=True)
        model_seed = args.seed + MODEL_SPECS.index(spec)
        torch.manual_seed(model_seed)
        torch.cuda.manual_seed_all(model_seed)
        print(f"  Sampling seed: {model_seed}", flush=True)
        rows: list[dict[str, object]] = []
        try:
            for case_index, (case_id, line_number, messages) in enumerate(cases, start=1):
                messages_for_generation = (
                    apply_age_gender_constraint(messages)
                    if spec.apply_age_gender_constraint
                    else messages
                )
                started = time.perf_counter()
                output, output_tokens, hit_token_limit = generate(
                    model,
                    tokenizer,
                    messages_for_generation,
                    args.max_new_tokens,
                    args.temperature,
                    args.top_p,
                )
                elapsed = round(time.perf_counter() - started, 3)
                rows.append(
                    {
                        "case_id": case_id,
                        "line_number": line_number,
                        "input": render_input(messages_for_generation),
                        "output": output,
                        "output_tokens": output_tokens,
                        "latency_seconds": elapsed,
                        "hit_token_limit": hit_token_limit,
                    }
                )
                print(
                    f"  case {case_index}/{len(cases)} complete: {case_id} ({elapsed:.3f}s)",
                    flush=True,
                )
        finally:
            del model
            del tokenizer
            gc.collect()
            torch.cuda.empty_cache()

        existing_sheet_index = None
        if spec.sheet_name in workbook.sheetnames:
            existing_sheet_index = workbook.index(workbook[spec.sheet_name])
            del workbook[spec.sheet_name]
        add_model_sheet(workbook, spec, resolved_adapter, rows, index=existing_sheet_index)
        if args.resume:
            add_resume_sampling_note(workbook, spec, args)
        workbook.save(args.output)
        print(f"  Workbook checkpoint saved: {args.output}", flush=True)

    verify_workbook(
        args.output,
        args.cases_per_model,
        MODEL_SPECS if args.resume else selected_specs,
    )
    print(f"Completed and verified: {args.output}", flush=True)


if __name__ == "__main__":
    main()
