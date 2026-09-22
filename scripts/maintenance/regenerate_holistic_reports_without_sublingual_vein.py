"""Regenerate four holistic-report cases with a system-and-user instruction.

The source workbook remains untouched.  A copied workbook receives new prompts and
outputs only in its ``全科报告`` worksheet; the other three model worksheets stay
as they were in the original batch run.
"""

from __future__ import annotations

import argparse
import gc
import shutil
import time
from pathlib import Path
from typing import Any

import torch
from infer_four_finetuned_models_to_excel import (
    DEFAULT_BASE_MODEL,
    generate,
    load_adapter,
    render_input,
    select_jsonl_validation_cases,
    style_sheet,
)
from openpyxl import load_workbook

HOLISTIC_ADAPTER = Path("artifacts/qwen3-4b-holistic-50k-qlora-windows")
HOLISTIC_VALIDATION = Path("data/holistic_50k/val.jsonl")
SOURCE_WORKBOOK = Path(
    "outputs/four_finetuned_models_4cases_20260827/four_finetuned_models_4cases.xlsx"
)
DEFAULT_OUTPUT = Path(
    "outputs/four_finetuned_models_4cases_20260827/"
    "four_finetuned_models_4cases_system_user_sublingual_vein_override.xlsx"
)
SUBLINGUAL_VEIN_INSTRUCTION = (
    "舌下脉络淡紫色属于正常范围，不得作为血瘀、血行不畅、寒凝或气滞的依据；"
    "相关判断只能依据其余四诊信息。"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Regenerate the holistic worksheet with an added sublingual-vein instruction."
    )
    parser.add_argument("--base-model", type=Path, default=DEFAULT_BASE_MODEL)
    parser.add_argument("--adapter", type=Path, default=HOLISTIC_ADAPTER)
    parser.add_argument("--source-workbook", type=Path, default=SOURCE_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--max-new-tokens", type=int, default=768)
    return parser.parse_args()


def add_instruction(messages: list[dict[str, str]]) -> list[dict[str, str]]:
    """Append the requested instruction to both system and final user messages."""
    updated = [dict(message) for message in messages]
    if len(updated) < 2 or updated[-1]["role"] != "assistant":
        raise ValueError("Expected complete SFT messages ending with an assistant reference")
    system_messages = [message for message in updated[:-1] if message["role"] == "system"]
    if not system_messages:
        raise ValueError("Expected at least one system message before the assistant reference")
    for system_message in system_messages:
        system_message["content"] = (
            f"{system_message['content'].rstrip()}\n\n{SUBLINGUAL_VEIN_INSTRUCTION}"
        )
    user_message = updated[-2]
    if user_message["role"] != "user":
        raise ValueError("Expected the assistant reference to be preceded by a user prompt")
    user_message["content"] = f"{user_message['content'].rstrip()}\n\n{SUBLINGUAL_VEIN_INSTRUCTION}"
    return updated


def update_holistic_sheet(workbook: Any, rows: list[dict[str, object]]) -> None:
    if "全科报告" not in workbook.sheetnames:
        raise ValueError("Source workbook is missing the 全科报告 worksheet")
    sheet = workbook["全科报告"]
    if sheet.max_row != 5 or sheet.max_column != 8:
        raise ValueError("Source workbook does not have the expected four-case 全科报告 layout")

    for row_index, row in enumerate(rows, start=2):
        sheet.cell(row_index, 1).value = "全科中医报告模型"
        sheet.cell(row_index, 2).value = row["case_id"]
        sheet.cell(row_index, 3).value = "data/holistic_50k/val.jsonl（验证集，system 与 user 均附加覆盖规则）"
        sheet.cell(row_index, 4).value = row["input"]
        sheet.cell(row_index, 5).value = row["output"]
        sheet.cell(row_index, 6).value = row["output_tokens"]
        sheet.cell(row_index, 7).value = row["latency_seconds"]
        sheet.cell(row_index, 8).value = "是" if row["hit_token_limit"] else "否"
        sheet.row_dimensions[row_index].height = 420


def add_overview_note(workbook: Any) -> None:
    overview = workbook["说明"]
    overview.append(["全科报告 system/user 覆盖规则", SUBLINGUAL_VEIN_INSTRUCTION])
    style_sheet(overview, (24, 108), 38)


def main() -> None:
    args = parse_args()
    if args.max_new_tokens <= 0:
        raise ValueError("--max-new-tokens must be positive")
    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is unavailable; QLoRA inference requires a CUDA GPU.")
    if not args.base_model.is_dir():
        raise FileNotFoundError(f"Base model not found: {args.base_model}")
    if not args.adapter.is_dir():
        raise FileNotFoundError(f"Adapter directory not found: {args.adapter}")
    if not args.source_workbook.is_file():
        raise FileNotFoundError(f"Source workbook not found: {args.source_workbook}")
    if args.output.exists():
        raise FileExistsError(f"Output already exists: {args.output}")

    selected_cases = select_jsonl_validation_cases(HOLISTIC_VALIDATION, 4)
    print("Loading holistic-report adapter...", flush=True)
    model, tokenizer, resolved_adapter = load_adapter(args.base_model, args.adapter)
    print(f"Adapter resolved to: {resolved_adapter}", flush=True)
    result_rows: list[dict[str, object]] = []
    try:
        for index, (case_id, original_messages) in enumerate(selected_cases, start=1):
            messages = add_instruction(original_messages)
            started = time.perf_counter()
            output, output_tokens, hit_token_limit = generate(
                model, tokenizer, messages, args.max_new_tokens
            )
            elapsed = round(time.perf_counter() - started, 3)
            result_rows.append(
                {
                    "case_id": case_id,
                    "input": render_input(messages),
                    "output": output,
                    "output_tokens": output_tokens,
                    "latency_seconds": elapsed,
                    "hit_token_limit": hit_token_limit,
                }
            )
            print(f"Case {index}/4 complete: {case_id} ({elapsed:.3f}s)", flush=True)
    finally:
        del model
        del tokenizer
        gc.collect()
        torch.cuda.empty_cache()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source_workbook, args.output)
    workbook = load_workbook(args.output)
    update_holistic_sheet(workbook, result_rows)
    add_overview_note(workbook)
    workbook.save(args.output)
    print(f"Completed: {args.output}", flush=True)


if __name__ == "__main__":
    main()
